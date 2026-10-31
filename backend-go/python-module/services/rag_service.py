"""Сервис для работы с RAG: загрузка документов, поиск и реранкинг."""

from __future__ import annotations

import logging
import os
from typing import Any

from adapters.model_manager import model_manager
from adapters.qdrant_repository import QdrantRepository
from core.exceptions import RAGError


class DocumentProcessor:
    """Обрабатывает документы в чанки для индексации."""

    def __init__(
        self,
        chunk_size: int = 512,
        chunk_overlap: int = 50,
    ) -> None:
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    def chunk_text(self, text: str) -> list[str]:
        if not text or not text.strip():
            return []
        words = text.split()
        chunks = []
        start = 0
        while start < len(words):
            end = min(start + self.chunk_size, len(words))
            chunk = " ".join(words[start:end])
            chunks.append(chunk)
            start += self.chunk_size - self.chunk_overlap
        return chunks

    def process_document(
        self, document: str, metadata: dict | None = None
    ) -> list[dict]:
        path = metadata.get("source") if metadata else None
        ext = path.lower().split(".")[-1] if path else None

        if ext == "pdf":
            content = self._extract_text_from_pdf(path)
        elif ext in ("ppt", "pptx"):
            content = self._extract_text_from_pptx(path)
        else:
            content = document

        chunks = self.chunk_text(content)
        return [
            {"chunk_id": i, "text": ch, "metadata": metadata or {}}
            for i, ch in enumerate(chunks)
        ]

    def _extract_text_from_pdf(self, path: str) -> str:
        import fitz

        if not fitz:
            raise RuntimeError("PyMuPDF (fitz) required. pip install pymupdf")
        text = []
        with fitz.open(path) as doc:
            for page in doc:
                text.append(page.get_text("text"))
        return "\n".join(text)

    def _extract_text_from_pptx(self, path: str) -> str:
        from pptx import Presentation

        if not Presentation:
            raise RuntimeError("python-pptx required. pip install python-pptx")
        prs = Presentation(path)
        text_runs = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_runs.append(shape.text)
        return "\n".join(text_runs)


class RAGSearcher:
    """Поиск с реранкингом через cross-encoder."""

    def __init__(
        self,
        qdrant_repo: QdrantRepository,
        cross_encoder_model: str,
    ) -> None:
        self.qdrant_repo = qdrant_repo
        self.cross_encoder = model_manager.get_cross_encoder_model(
            cross_encoder_model
        )

    def search_raw_candidates(
        self, query: str, top_k: int = 30
    ) -> list[dict[str, Any]]:
        query_vector = self.qdrant_repo.encoder.encode([query])[0].tolist()
        return self.qdrant_repo.search(query_vector, top_k=top_k)

    def rerank_results(
        self, query: str, candidates: list[dict[str, Any]], limit: int = 10
    ) -> list[dict[str, Any]]:
        pairs = [(query, c.get("text") or "") for c in candidates]
        scores = self.cross_encoder.predict(pairs, batch_size=16)

        for i, score in enumerate(scores):
            candidates[i]["rerank_score"] = float(score)

        sorted_candidates = sorted(
            candidates, key=lambda x: x["rerank_score"], reverse=True
        )
        return sorted_candidates[:limit]

    def search(
        self, query: str, top_k: int = 30, rerank_limit: int = 5
    ) -> list[dict[str, Any]]:
        candidates = self.search_raw_candidates(query, top_k=top_k)
        reranked = self.rerank_results(query, candidates, limit=rerank_limit)
        return reranked


class RAGService:
    """Оркестрирует загрузку документов в Qdrant и семантический поиск."""

    def __init__(
        self,
        qdrant_repo: QdrantRepository,
        collection_name: str = "atomic_hack",
        cross_encoder_model: str | None = None,
        chunk_size: int = 512,
        chunk_overlap: int = 50,
    ) -> None:
        self.qdrant_repo = qdrant_repo
        self.collection_name = collection_name
        self.doc_processor = DocumentProcessor(
            chunk_size=chunk_size, chunk_overlap=chunk_overlap
        )
        self._searcher: RAGSearcher | None = None
        self._cross_encoder_model = cross_encoder_model
        self.logger = logging.getLogger(__name__)

    def _get_searcher(self) -> RAGSearcher:
        if self._searcher is None:
            if not self._cross_encoder_model:
                from config import settings

                self._cross_encoder_model = settings.CROSS_ENCODER_MODEL
            self._searcher = RAGSearcher(
                self.qdrant_repo, self._cross_encoder_model
            )
        return self._searcher

    def load_documents(
        self,
        documents: list[str],
        metadata: list[dict[str, Any]] | None = None,
    ) -> int:
        all_chunks = []
        metadata = metadata or []
        for i, doc in enumerate(documents):
            meta = metadata[i] if i < len(metadata) else {"doc_id": i}
            path = meta.get("source")
            if path and os.path.exists(path):
                chunks = self.doc_processor.process_document(doc, meta)
            else:
                chunks = self.doc_processor.process_document(doc, meta)
            all_chunks.extend(chunks)

        self.qdrant_repo.add_documents(all_chunks)
        self.logger.info("Loaded %d chunks into Qdrant", len(all_chunks))
        return len(all_chunks)

    def search(
        self,
        query: str,
        top_k: int = 30,
        rerank_limit: int = 5,
    ) -> list[dict[str, Any]]:
        searcher = self._get_searcher()
        return searcher.search(query, top_k=top_k, rerank_limit=rerank_limit)

    def rerank(
        self,
        query: str,
        candidates: list[dict[str, Any]],
        limit: int = 10,
    ) -> list[dict[str, Any]]:
        searcher = self._get_searcher()
        return searcher.rerank_results(query, candidates, limit=limit)

    def get_searcher(self) -> RAGSearcher:
        return self._get_searcher()